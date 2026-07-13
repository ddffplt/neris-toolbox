#!/usr/bin/env python3
"""Build one combined report workbook per entity from prompted NERIS outputs."""

import re
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl

ENTITY_ID = re.compile(r"^FD[A-Za-z0-9]+$")


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._") or "department"


def prompt_files():
    folder = Path(input("Folder containing NERIS output Excel files: ").strip().strip('"'))
    if not folder.is_dir():
        raise ValueError(f"Folder not found: {folder}")
    files = sorted(folder.glob("*.xlsx"))
    if not files:
        raise ValueError(f"No .xlsx files found in {folder}")

    print("\nAvailable Excel files:")
    for index, path in enumerate(files, 1):
        print(f"  {index}. {path.name}")
    selection = input("Select files by number, separated by commas (Enter for all): ").strip()
    if selection:
        try:
            indexes = [int(value.strip()) for value in selection.split(",")]
            selected = [files[index - 1] for index in indexes]
        except (ValueError, IndexError):
            raise ValueError("Select valid file numbers separated by commas.")
    else:
        selected = files
    return selected


def prompt_output_folder(input_folder):
    default = input_folder / "department-reports"
    value = input(f"Output folder [{default}]: ").strip().strip('"')
    return Path(value) if value else default


def read_rows(path, sheet_name):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        return None, None
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return None, None
    headers = list(rows[0])
    entity_rows = [
        list(row) for row in rows[1:]
        if row and row[0] and ENTITY_ID.match(str(row[0]))
    ]
    return headers, entity_rows


def read_fire_blocks(path):
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    worksheet = workbook["Entity Fire Reports"]
    blocks = {}
    current_id = None
    current_name = ""
    current_start = None
    for row_number, row in enumerate(worksheet.iter_rows(), start=1):
        first = row[0].value
        second = row[1].value if len(row) > 1 else None
        if first and second and ENTITY_ID.match(str(first)):
            if current_id:
                blocks[current_id] = (current_name, current_start, row_number - 1)
            current_id = str(first)
            current_name = str(second)
            current_start = row_number
    if current_id:
        blocks[current_id] = (current_name, current_start, worksheet.max_row)
    return workbook, worksheet, blocks


def add_fire_sheet(workbook, source_worksheet, block):
    worksheet = workbook.create_sheet("Fire Activity")
    _, start_row, end_row = block
    offset = 1 - start_row
    for source_row in range(start_row, end_row + 1):
        target_row = source_row + offset
        if source_row in source_worksheet.row_dimensions:
            worksheet.row_dimensions[target_row].height = source_worksheet.row_dimensions[source_row].height
        for source_column in range(1, source_worksheet.max_column + 1):
            source_cell = source_worksheet.cell(source_row, source_column)
            target_cell = worksheet.cell(target_row, source_column, source_cell.value)
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.number_format = source_cell.number_format
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
    for letter, dimension in source_worksheet.column_dimensions.items():
        worksheet.column_dimensions[letter].width = dimension.width
        worksheet.column_dimensions[letter].hidden = dimension.hidden
    for merged in source_worksheet.merged_cells.ranges:
        if merged.min_row >= start_row and merged.max_row <= end_row:
            worksheet.merge_cells(
                start_row=merged.min_row + offset,
                start_column=merged.min_col,
                end_row=merged.max_row + offset,
                end_column=merged.max_col,
            )
    worksheet.freeze_panes = source_worksheet.freeze_panes
    worksheet.sheet_view.showGridLines = source_worksheet.sheet_view.showGridLines


def add_table(workbook, title, headers, rows, entity_name, entity_id):
    worksheet = workbook.create_sheet(title[:31])
    worksheet.append([entity_name, entity_id])
    worksheet.append([])
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A4"
    for cell in worksheet[3]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="1F4E78")
    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2, 36
        )


def add_sources_sheet(workbook, entity_name, entity_id, source_files):
    worksheet = workbook.create_sheet("Report Sources")
    worksheet.append([entity_name, entity_id])
    worksheet.append(["Generated at", datetime.now().astimezone().isoformat(timespec="seconds")])
    worksheet.append(["Source file", "Selected input file"])
    for category, path in source_files:
        worksheet.append([category, path.name])
    worksheet.append(["Entity isolation", "Only rows belonging to the entity shown above were included."])
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 90
    for cell in worksheet[3]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="1F4E78")


def classify_files(files):
    classified = {}
    for path in files:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = set(workbook.sheetnames)
        if "Entity Fire Reports" in sheets:
            category = "fire"
        elif "Incident Type Summary" in sheets:
            category = "types"
        elif "Monthly Incident Counts" in sheets:
            category = "monthly"
        else:
            print(f"Skipping unrecognized workbook: {path.name}")
            continue
        if category in classified:
            raise ValueError(f"More than one {category} report selected; select only one of each type.")
        classified[category] = path
    if "fire" not in classified:
        raise ValueError("Select a fire activity workbook containing 'Entity Fire Reports'.")
    return classified


def build_reports(files, output_folder):
    classified = classify_files(files)
    fire_path = classified["fire"]
    fire_workbook, fire_worksheet, fire_blocks = read_fire_blocks(fire_path)
    type_headers, type_rows = read_rows(classified["types"], "Incident Type Summary") if "types" in classified else (None, [])
    type_month_headers, type_month_rows = read_rows(classified["types"], "Monthly Incident Counts") if "types" in classified else (None, [])
    month_headers, month_rows = read_rows(classified["monthly"], "Monthly Incident Counts") if "monthly" in classified else (None, [])
    rows_by_entity = {
        "types": {str(row[0]): row for row in type_rows},
        "type_month": {str(row[0]): row for row in type_month_rows},
        "monthly": {str(row[0]): row for row in month_rows},
    }
    entity_ids = sorted(set(fire_blocks) | set(rows_by_entity["types"]) | set(rows_by_entity["monthly"]))
    output_folder.mkdir(parents=True, exist_ok=True)

    for entity_id in entity_ids:
        entity_name = fire_blocks.get(entity_id, ("", None, None))[0]
        entity_name = entity_name or rows_by_entity["types"].get(entity_id, ["", entity_id])[1] or entity_id
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        fire_block = fire_blocks.get(entity_id)
        if fire_block:
            add_fire_sheet(workbook, fire_worksheet, fire_block)
        if "types" in classified and entity_id in rows_by_entity["types"]:
            add_table(workbook, "Incident Types", type_headers, [rows_by_entity["types"][entity_id]], entity_name, entity_id)
        if "types" in classified and entity_id in rows_by_entity["type_month"]:
            add_table(workbook, "Type Monthly Activity", type_month_headers, [rows_by_entity["type_month"][entity_id]], entity_name, entity_id)
        if "monthly" in classified and entity_id in rows_by_entity["monthly"]:
            add_table(workbook, "State Monthly Activity", month_headers, [rows_by_entity["monthly"][entity_id]], entity_name, entity_id)
        add_sources_sheet(workbook, entity_name, entity_id, sorted(classified.items()))
        path = output_folder / f"{safe_name(entity_name)}_{safe_name(entity_id)}_combined.xlsx"
        workbook.save(path)
        print(f"Created: {path}")
    print(f"Created {len(entity_ids)} combined report(s).")


def main():
    print("\nNERIS Department Reports by Prompt")
    files = prompt_files()
    output_folder = prompt_output_folder(files[0].parent)
    build_reports(files, output_folder)


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError) as error:
        raise SystemExit(str(error))
