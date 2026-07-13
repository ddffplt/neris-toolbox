#!/usr/bin/env python3
"""Create and email one NERIS activity workbook per fire department."""

import argparse
import importlib.util
import os
import re
import smtplib
import sys
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import EntityFireActivityReport as fire_report
import StateActivityReport_w_IncTypeCount as state_report


SCRIPT_DIR = Path(__file__).resolve().parent
CLIENT_STATE_REPORT = SCRIPT_DIR / "client-credentials" / "stateactivity.py"
SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")


def load_client_state_report():
    """Load the client-credential report without changing its module name."""
    spec = importlib.util.spec_from_file_location("client_state_report", CLIENT_STATE_REPORT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--state", required=True, help="Two-letter state code")
    parser.add_argument("--entity", help="Only process this NERIS entity ID")
    parser.add_argument("--start", help="UTC start date, YYYY-MM-DD")
    parser.add_argument("--end", help="UTC end date, YYYY-MM-DD")
    parser.add_argument("--output-dir", default="department-reports")
    parser.add_argument("--limit", type=int, help="Only process the first N departments")
    parser.add_argument("--send", action="store_true", help="Send email; otherwise only create files")
    return parser.parse_args()


def get_date_range(args):
    start = datetime.strptime(args.start, "%Y-%m-%d") if args.start else None
    end = datetime.strptime(args.end, "%Y-%m-%d") if args.end else None
    if start and end and end < start:
        raise ValueError("--end must be on or after --start")
    return state_report.calculate_date_range(
        "Custom Date Range" if start or end else "All Records",
        start.date() if start else None,
        end.date() if end else None,
    )


def get_department_emails(client, entity_id):
    """Return active, non-empty emails for users belonging to one department."""
    base_url = os.environ.get("NERIS_BASE_URL", "https://api.neris.fsri.org/v1")
    response = client._session.get(
        f"{base_url}/entity/{entity_id}/user_entity_membership",
        params={"page_size": 100, "page_number": 1},
    )
    data = response if isinstance(response, dict) else response.json()
    emails = set()
    for user in data.get("users", []):
        email = (user.get("email") or "").strip()
        if email and user.get("active", True):
            emails.add(email)
    return sorted(emails)


def safe_filename(value):
    return SAFE_NAME.sub("_", value).strip("._") or "department"


def build_workbook(path, entity_id, entity_name, fire_data, state_data,
                   client_state_data, month_columns, range_label):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    def add_table(title, headers, rows):
        sheet = workbook.create_sheet(title[:31])
        sheet.append([entity_name, entity_id])
        sheet.append(["Date range", range_label])
        sheet.append([])
        sheet.append(headers)
        for cell in sheet[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A5"
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)

    sources = workbook.create_sheet("Report Sources")
    sources.append([entity_name, entity_id])
    sources.append(["Report date/range", range_label])
    sources.append(["Generated at", datetime.now().astimezone().isoformat(timespec="seconds")])
    sources.append([])
    sources.append(["Report source", "Data included"])
    sources.append(["EntityFireActivityReport.py", "Fire activity metrics and fire incident summaries"])
    sources.append(["StateActivityReport_w_IncTypeCount.py", "Incident type counts, monthly activity, and NAR status"])
    sources.append(["client-credentials/stateactivity.py", "Client-credential monthly incident activity cross-check"])
    sources.column_dimensions["A"].width = 48
    sources.column_dimensions["B"].width = 85
    for cell in sources[5]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    detail, summary = fire_data
    add_table(
        "Fire Activity",
        ["Fire Incident Type"] + fire_report.METRIC_COLUMNS,
        [[type_name] + [values[column] for column in fire_report.METRIC_COLUMNS]
         for (eid, type_name), values in sorted(detail.items()) if eid == entity_id],
    )
    add_table(
        "Fire Summaries",
        ["Fire Summary Type"] + fire_report.METRIC_COLUMNS,
        [[type_name] + [values[column] for column in fire_report.METRIC_COLUMNS]
         for (eid, type_name), values in sorted(summary.items()) if eid == entity_id],
    )
    type_values = state_data.get("types", {})
    add_table(
        "Incident Types",
        ["Incident Type", "Count"],
        [[name, count] for name, count in sorted(type_values.items())],
    )
    add_table(
        "Monthly Activity",
        ["Month", "Incident Count", "NAR"],
        [[month, state_data.get("months", {}).get(month, ""),
          "Yes" if month in state_data.get("nar_months", set()) else ""]
         for month in month_columns],
    )
    add_table(
        "Client Credential Check",
        ["Month", "Incident Count"],
        [[month, client_state_data.get(month, "")] for month in month_columns],
    )
    workbook.save(path)


def send_report(sender, recipients, subject, body, attachment):
    message = EmailMessage()
    message["From"] = sender
    message["To"] = ", ".join(recipients)
    message["Subject"] = subject
    message.set_content(body)
    with attachment.open("rb") as stream:
        message.add_attachment(
            stream.read(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=attachment.name,
        )
    host = os.environ["SMTP_HOST"]
    port = int(os.environ.get("SMTP_PORT", "587"))
    with smtplib.SMTP(host, port) as smtp:
        smtp.starttls()
        smtp.login(os.environ["SMTP_USERNAME"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(message)


def main():
    args = parse_args()
    if len(args.state) != 2:
        raise ValueError("--state must be a two-letter state code")
    start, end = get_date_range(args)
    range_label = f"{args.start or 'all'}_to_{args.end or 'present'}"

    fire_report.ensure_dependencies()
    username = input("NERIS Email: ").strip()
    password = input("NERIS Password: ").strip()
    client = fire_report.authenticate(username, password)
    entities = fire_report.fetch_entities(client, args.state, args.entity)
    if args.limit is not None:
        if args.limit < 1:
            raise ValueError("--limit must be at least 1")
        entities = dict(sorted(entities.items())[:args.limit])
        print(f"Testing first {len(entities)} departments; email sending remains disabled unless --send is supplied.")
    client_state_report = load_client_state_report()
    month_columns = state_report.generate_month_columns(start, end)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    smtp_sender = os.environ.get("SMTP_FROM", os.environ.get("SMTP_USERNAME", ""))
    for entity_id, entity_name in sorted(entities.items()):
        print(f"Processing {entity_name or entity_id}...")
        incidents = fire_report.fetch_incidents(client, args.state, entity_id, start, end)
        detail, _, summary, _, _ = fire_report.aggregate(incidents, {entity_id: entity_name})
        state_data = state_report.fetch_incidents_for_entity(
            client, args.state, entity_id, start, end
        )
        state_data["nar_months"] = {
            month for month in state_report.fetch_nars_for_entity(
                client, args.state, entity_id
            ) if month in month_columns
        }
        client_counts = client_state_report.fetch_incidents_for_entity(client, args.state, entity_id)
        client_counts = {
            month: count for month, count in client_counts.items()
            if month in month_columns
        }
        filename = f"{safe_filename(entity_name)}_{safe_filename(entity_id)}_{range_label}.xlsx"
        path = output_dir / filename
        build_workbook(path, entity_id, entity_name, (detail, summary), state_data,
                       client_counts, month_columns, range_label)
        recipients = get_department_emails(client, entity_id)
        if args.send:
            if not smtp_sender:
                raise ValueError("Set SMTP_FROM or SMTP_USERNAME before using --send")
            if not recipients:
                print("  No active department email addresses; attachment retained.")
                continue
            send_report(
                smtp_sender, recipients,
                f"NERIS activity report: {entity_name or entity_id}",
                f"Attached is the NERIS activity report for {entity_name or entity_id}.\n\n"
                f"Reporting period: {range_label}",
                path,
            )
            print(f"  Sent to {', '.join(recipients)}")
        else:
            print(f"  Created {path} (dry run; not sent)")


if __name__ == "__main__":
    try:
        main()
    except (KeyError, ValueError) as error:
        sys.exit(str(error))