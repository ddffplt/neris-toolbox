#!/usr/bin/env python3
# pyright: reportMissingImports=false
"""NERIS firefighter and civilian casualty report for fire incidents, by month."""

import os
import sys
from collections import defaultdict
from datetime import datetime, timezone

from EntityFireActivityReport import (
    DATE_RANGE_OPTIONS as _BASE_DATE_RANGE_OPTIONS,
    _metric_values,
    _type_strings,
    calculate_date_range as _base_calculate_date_range,
    date_range_label as _base_date_range_label,
    ensure_dependencies,
    fetch_entities,
    fetch_incidents,
)

CURRENT_YTD = "Occurred in current calendar year to date"
DATE_RANGE_OPTIONS = _BASE_DATE_RANGE_OPTIONS[:-1] + [CURRENT_YTD, _BASE_DATE_RANGE_OPTIONS[-1]]


def calculate_date_range(range_type, start_date=None, end_date=None):
    if range_type == CURRENT_YTD:
        now = datetime.now(timezone.utc)
        return datetime(now.year, 1, 1, tzinfo=timezone.utc), now
    return _base_calculate_date_range(range_type, start_date, end_date)


def date_range_label(range_type, start_date=None, end_date=None):
    if range_type == CURRENT_YTD:
        return f"{datetime.now().year}-ytd"
    return _base_date_range_label(range_type, start_date, end_date)


CASUALTY_COLUMNS = [
    "Number of Fires",
    "Number of FF Deaths", "Number of FF Injuries",
    "Number of nonFF Deaths", "Number of nonFF Injuries",
    "Number of FF Rescues", "Number of nonFF Rescues",
]


def authenticate_client_credentials():
    """Authenticate with an OAuth client; no email, password, or MFA needed."""
    from neris_api_client import Config, NerisApiClient
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    client_id = "ffcc03c8-7a39-448c-b802-307912a1a29a"
    client_secret = "135q1hpi066ZnCyXXE5b2g"

    print("\nConnecting to NERIS API...")
    client = NerisApiClient(
        Config(
            base_url="https://api.neris.fsri.org/v1",
            grant_type="client_credentials",
            client_id=client_id,
            client_secret=client_secret,
        )
    )
    client._session.verify = False
    client.list_incidents(page_size=1)
    print("✓ Authentication successful!")
    return client


def prompt_config():
    print("\n" + "=" * 60)
    print("  NERIS Fire Casualty Report")
    print("=" * 60)
    state_code = input("State Code (e.g. VA, MI, CA): ").strip().upper()

    print("\nScope Options:")
    print("  1. Entire state")
    print("  2. County (FD25 prefix)")
    print("  3. Single entity")
    while True:
        scope = input("Select scope [1]: ").strip() or "1"
        if scope in {"1", "2", "3"}:
            break
        print("  Invalid selection, choose 1, 2, or 3.")

    entity_id = entity_prefix = None
    if scope == "2":
        while True:
            entity_prefix = input("County prefix (e.g. FD25001): ").strip().upper().rstrip("*")
            if entity_prefix.startswith("FD") and len(entity_prefix) >= 6:
                break
            print("  Enter a prefix such as FD25001.")
    elif scope == "3":
        while True:
            entity_id = input("NERIS Entity ID (e.g. FD25001001): ").strip().upper()
            if entity_id:
                break
            print("  An entity ID is required.")

    print("\nDate Range Options:")
    for number, option in enumerate(DATE_RANGE_OPTIONS, 1):
        print(f"  {number}. {option}")
    while True:
        choice = input("Select option [1]: ").strip() or "1"
        if choice.isdigit() and 1 <= int(choice) <= len(DATE_RANGE_OPTIONS):
            range_type = DATE_RANGE_OPTIONS[int(choice) - 1]
            break
        print("  Invalid selection, try again.")

    start_date = end_date = None
    if range_type == "Custom Date Range":
        while True:
            try:
                start_date = datetime.strptime(input("Start date (YYYY-MM-DD): ").strip(), "%Y-%m-%d").date()
                break
            except ValueError:
                print("  Invalid format, use YYYY-MM-DD.")
        while True:
            try:
                end_date = datetime.strptime(input("End date (YYYY-MM-DD): ").strip(), "%Y-%m-%d").date()
                if end_date >= start_date:
                    break
                print("  End date must be on or after the start date.")
            except ValueError:
                print("  Invalid format, use YYYY-MM-DD.")

    if not state_code:
        sys.exit("State code is required.")
    return state_code, scope, entity_id, entity_prefix, range_type, start_date, end_date


def incident_month(incident):
    value = (incident.get("dispatch") or {}).get("call_create") or ""
    return str(value)[:7] if len(str(value)) >= 7 else None


def aggregate_by_month(incidents):
    months = defaultdict(lambda: {column: 0 for column in CASUALTY_COLUMNS})
    for incident in incidents:
        if not _type_strings(incident):
            continue
        month = incident_month(incident)
        if not month:
            continue
        metrics = _metric_values(incident)
        months[month]["Number of Fires"] += 1
        for column in CASUALTY_COLUMNS[1:]:
            months[month][column] += metrics[column]
    return months


def aggregate_by_entity(incidents):
    entities = defaultdict(lambda: {column: 0 for column in CASUALTY_COLUMNS})
    for incident in incidents:
        if not _type_strings(incident):
            continue
        entity_id = (incident.get("base") or {}).get("department_neris_id", "")
        if not entity_id:
            continue
        metrics = _metric_values(incident)
        entities[entity_id]["Number of Fires"] += 1
        for column in CASUALTY_COLUMNS[1:]:
            entities[entity_id][column] += metrics[column]
    return entities


def month_sequence(months):
    if not months:
        return []
    first, last = min(months), max(months)
    year, month = int(first[:4]), int(first[5:7])
    result = []
    while f"{year:04d}-{month:02d}" <= last:
        result.append(f"{year:04d}-{month:02d}")
        month += 1
        if month > 12:
            year, month = year + 1, 1
    return result


def month_label(month):
    year, month_number = int(month[:4]), int(month[5:7])
    return f"{year} - {datetime(year, month_number, 1).strftime('%b')}"


def export_excel(months, entities, entity_names, scope_label, state_code, range_label, range_type):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(
        os.path.expanduser("~"), "Documents",
        f"neris_fire_casualties_{state_code}_{range_label}_{timestamp}.xlsx",
    )
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    side = Side(style="thin", color="B7C9D6")
    border = Border(left=side, right=side, top=side, bottom=side)

    def style_header(sheet, row_number, count):
        for column_index in range(1, count + 1):
            cell = sheet.cell(row_number, column_index)
            cell.fill, cell.font, cell.border = header_fill, header_font, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Entity Summary sheet: one row per entity with casualty totals for the date range.
    entity_sheet = workbook.active
    entity_sheet.title = "Entity Summary"
    entity_sheet.append([f"Fire casualty summary by entity: {scope_label}"])
    entity_sheet.cell(entity_sheet.max_row, 1).font = Font(bold=True, size=12)
    entity_headers = ["Entity ID", "Entity Name"] + CASUALTY_COLUMNS
    entity_sheet.append(entity_headers)
    header_row = entity_sheet.max_row
    style_header(entity_sheet, header_row, len(entity_headers))
    entity_sheet.row_dimensions[header_row].height = 32

    entity_totals = {column: 0 for column in CASUALTY_COLUMNS}
    for entity_id in sorted(entity_names):
        values = entities.get(entity_id, {column: 0 for column in CASUALTY_COLUMNS})
        entity_sheet.append(
            [entity_id, entity_names[entity_id]] + [values[column] for column in CASUALTY_COLUMNS]
        )
        for column in CASUALTY_COLUMNS:
            entity_totals[column] += values[column]
        for column_index in range(1, len(entity_headers) + 1):
            entity_sheet.cell(entity_sheet.max_row, column_index).border = border
    entity_sheet.append(["Total", ""] + [entity_totals[column] for column in CASUALTY_COLUMNS])
    for column_index in range(1, len(entity_headers) + 1):
        cell = entity_sheet.cell(entity_sheet.max_row, column_index)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.border = border

    entity_sheet.column_dimensions["A"].width = 16
    entity_sheet.column_dimensions["B"].width = 40
    for column_index in range(3, len(entity_headers) + 1):
        entity_sheet.column_dimensions[get_column_letter(column_index)].width = 14

    # Monthly Breakdown sheet: casualty totals aggregated across all selected entities, by month.
    month_sheet = workbook.create_sheet("Monthly Breakdown")
    month_sheet.append([f"Monthly breakdown: {scope_label}"])
    month_sheet.cell(month_sheet.max_row, 1).font = Font(bold=True, size=12)
    headers = ["Month"] + CASUALTY_COLUMNS
    month_sheet.append(headers)
    header_row = month_sheet.max_row
    style_header(month_sheet, header_row, len(headers))
    month_sheet.row_dimensions[header_row].height = 32

    totals = {column: 0 for column in CASUALTY_COLUMNS}
    for month in month_sequence(months):
        values = months.get(month, {column: 0 for column in CASUALTY_COLUMNS})
        month_sheet.append([month_label(month)] + [values[column] for column in CASUALTY_COLUMNS])
        for column in CASUALTY_COLUMNS:
            totals[column] += values[column]
        for column_index in range(1, len(headers) + 1):
            month_sheet.cell(month_sheet.max_row, column_index).border = border
    month_sheet.append(["Total"] + [totals[column] for column in CASUALTY_COLUMNS])
    for column_index in range(1, len(headers) + 1):
        cell = month_sheet.cell(month_sheet.max_row, column_index)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
        cell.border = border

    month_sheet.column_dimensions["A"].width = 16
    for column_index in range(2, len(headers) + 1):
        month_sheet.column_dimensions[get_column_letter(column_index)].width = 14

    from report_metadata import add_report_sources_sheet

    add_report_sources_sheet(
        workbook, "FireCasualtyReport.py", range_label,
        [("State code", state_code), ("Scope", scope_label), ("Date range", range_type),
         ("Incident filter", "Incidents with a FIRE incident type")],
    )
    workbook.save(path)
    return path


def main():
    ensure_dependencies()
    (state_code, scope, entity_id, entity_prefix,
     range_type, start_date, end_date) = prompt_config()
    start, end = calculate_date_range(range_type, start_date, end_date)
    client = authenticate_client_credentials()

    entity_names = fetch_entities(client, state_code, entity_id)
    if entity_prefix:
        entity_names = {eid: name for eid, name in entity_names.items()
                        if str(eid).upper().startswith(entity_prefix)}
    if not entity_names:
        print("No entities found for the selected scope.")
        return

    incidents = []
    failed = []
    for current_entity in entity_names:
        print(f"Fetching {current_entity}...")
        try:
            incidents.extend(fetch_incidents(client, state_code, current_entity, start, end))
        except RuntimeError as error:
            print(f"  FAILED {current_entity}: {error}")
            failed.append(current_entity)
    if failed:
        print(f"\nWARNING: {len(failed)} entities could not be retrieved and are missing from the report: "
              f"{', '.join(failed)}")

    scope_label = {"1": f"Entire state ({state_code})",
                   "2": f"County {entity_prefix}",
                   "3": f"Entity {entity_id}"}[scope]
    months = aggregate_by_month(incidents)
    entities = aggregate_by_entity(incidents)
    path = export_excel(months, entities, entity_names, scope_label, state_code,
                        date_range_label(range_type, start_date, end_date), range_type)
    print(f"\nExported: {path}")
    print(f"Incidents retrieved: {len(incidents)}")


if __name__ == "__main__":
    main()
