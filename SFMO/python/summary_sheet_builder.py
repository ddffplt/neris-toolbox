"""Reusable department summary metrics and Excel worksheet builder."""

from collections import defaultdict

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SUMMARY_HEADERS = [
    "Department/Entity ID", "Department Name", "Total Number of Incidents",
    "Firefighter Fatal Casualties", "Firefighter Nonfatal Casualties",
    "Civilian Fatal Casualties", "Civilian Nonfatal Casualties",
    "Total Dollar Loss", "Acres Burned",
]

METRIC_NAMES = [
    "Total Number of Incidents", "Firefighter Fatal Casualties",
    "Firefighter Nonfatal Casualties", "Civilian Fatal Casualties",
    "Civilian Nonfatal Casualties", "Total Dollar Loss", "Acres Burned",
]


def _number_value(value):
    if value is None or value == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").replace("$", ""))
    except (TypeError, ValueError):
        return 0.0


def _first_number(incident, paths):
    for path in paths:
        value = incident
        for part in path.split("."):
            if not isinstance(value, dict):
                value = None
                break
            value = value.get(part)
        if value is not None and value != "":
            return _number_value(value)
    return 0.0


def _is_fatal(injury_type):
    value = str(injury_type or "").upper()
    return "FATAL" in value and "NONFATAL" not in value


def _incident_metrics(incident):
    property_loss = _first_number(incident, [
        "fire_detail.property_value_loss",
        "fire_detail.property_value_loss_amount",
        "fire_detail.location_detail.property_value_loss",
    ])
    contents_loss = _first_number(incident, [
        "fire_detail.contents_value_loss",
        "fire_detail.contents_value_loss_amount",
        "fire_detail.location_detail.contents_value_loss",
    ])
    metrics = {
        "Total Number of Incidents": 1,
        "Firefighter Fatal Casualties": 0,
        "Firefighter Nonfatal Casualties": 0,
        "Civilian Fatal Casualties": 0,
        "Civilian Nonfatal Casualties": 0,
        "Total Dollar Loss": property_loss + contents_loss,
        "Acres Burned": _first_number(incident, [
            "fire_detail.outside_fire_acres_burned",
            "outside_fire_acres_burned",
            "fire_detail.location_detail.acres_burned",
        ]),
    }
    for casualty_rescue in incident.get("casualty_rescues") or []:
        if not isinstance(casualty_rescue, dict):
            continue
        casualty_type = str(casualty_rescue.get("type") or "").upper()
        injury = (casualty_rescue.get("casualty") or {}).get(
            "injury_or_noninjury"
        ) or {}
        injury_type = str(injury.get("type") or "").upper()
        if _is_fatal(injury_type):
            suffix = "Fatal"
        elif "NONFATAL" in injury_type:
            suffix = "Nonfatal"
        else:
            continue
        prefix = "Firefighter" if casualty_type == "FF" else "Civilian"
        metrics[f"{prefix} {suffix} Casualties"] += 1
    return metrics


def build_summary_metrics(incidents, department_names=None, entity_ids=None):
    """Aggregate requested summary metrics by department/entity ID.

    When entity_ids is supplied, every known entity is included in the summary,
    even if it has no matching incidents in the current report.
    """
    department_names = department_names or {}
    entity_ids = {str(entity_id) for entity_id in (entity_ids or []) if entity_id}
    summary_metrics = defaultdict(
        lambda: {"Department Name": "", **{name: 0 for name in METRIC_NAMES}}
    )
    for entity_id in entity_ids:
        summary_metrics[entity_id]["Department Name"] = department_names.get(entity_id, "")

    for incident in incidents:
        entity_id = str((incident.get("base") or {}).get("department_neris_id", ""))
        if not entity_id:
            continue
        summary_metrics[entity_id]["Department Name"] = department_names.get(entity_id, "")
        for metric, value in _incident_metrics(incident).items():
            summary_metrics[entity_id][metric] += value
    return dict(summary_metrics)


def write_summary_sheet(workbook, summary_metrics, entity_ids=None):
    """Add a formatted Summary worksheet to an openpyxl workbook."""
    worksheet = workbook.create_sheet("Summary")
    worksheet.append(SUMMARY_HEADERS)
    target_ids = {str(entity_id) for entity_id in (entity_ids or []) if entity_id}
    for entity_id in sorted(set(target_ids) | set(summary_metrics)):
        totals = summary_metrics.get(entity_id, {"Department Name": "", **{name: 0 for name in METRIC_NAMES}})
        worksheet.append([
            entity_id,
            totals.get("Department Name", ""),
            totals.get("Total Number of Incidents", 0),
            totals.get("Firefighter Fatal Casualties", 0),
            totals.get("Firefighter Nonfatal Casualties", 0),
            totals.get("Civilian Fatal Casualties", 0),
            totals.get("Civilian Nonfatal Casualties", 0),
            totals.get("Total Dollar Loss", 0),
            totals.get("Acres Burned", 0),
        ])

    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(color="FFFFFF", bold=True, size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    for row in worksheet.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.number_format = "#,##0;(#,##0);-"
    for cell in worksheet["H"][1:]:
        cell.number_format = "$#,##0;($#,##0);-"
    for cell in worksheet["I"][1:]:
        cell.number_format = "#,##0.00;(#,##0.00);-"
    for column_index, header in enumerate(SUMMARY_HEADERS, 1):
        column_letter = get_column_letter(column_index)
        max_length = len(header)
        for row in worksheet.iter_rows(min_col=column_index, max_col=column_index):
            if row[0].value is not None:
                max_length = max(max_length, len(str(row[0].value)))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)
    worksheet.freeze_panes = "A2"
    return worksheet
