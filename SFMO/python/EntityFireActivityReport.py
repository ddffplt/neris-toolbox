#!/usr/bin/env python3
"""NERIS fire activity report by entity and fire incident type."""

import os
import subprocess
import sys
import math
from collections import defaultdict
from datetime import datetime, timedelta, timezone

DATE_RANGE_OPTIONS = [
    "All Records", "Occurred Today", "Occurred Yesterday",
    "Occurred Last 7 Days", "Occurred Last 30 Days", "Occurred Last 90 Days",
    "Occurred in 2025", "Custom Date Range",
]

METRIC_COLUMNS = [
    "Number of Fires", "Number of Exposures",
    "Number of FF Deaths", "Number of FF Injuries",
    "Number of nonFF Deaths", "Number of nonFF Injuries",
    "Number of FF Rescues", "Number of nonFF Rescues",
    "Number of Displacements", "Aid Given In Support of",
    "Aid Given In Lieu of", "Aid Given Acting As",
    "Aid Received In Support of", "Aid Received In Lieu of",
    "Aid Received Acting As", "Property Value Loss",
    "Contents Value Loss", "Total Loss", "Acres Burned",
]

SUMMARY_GROUPS = {
    "Structure Fires": "STRUCTURE_FIRE",
    "Transportation Fires": "TRANSPORTATION_FIRE",
    "Special Fires": "SPECIAL_FIRE",
    "Outside Fires": "OUTSIDE_FIRE",
}

AID_COLUMNS = {
    ("GIVEN", "IN_SUPPORT_OF"): "Aid Given In Support of",
    ("GIVEN", "IN_LIEU_OF"): "Aid Given In Lieu of",
    ("GIVEN", "ACTING_AS"): "Aid Given Acting As",
    ("RECEIVED", "IN_SUPPORT_OF"): "Aid Received In Support of",
    ("RECEIVED", "IN_LIEU_OF"): "Aid Received In Lieu of",
    ("RECEIVED", "ACTING_AS"): "Aid Received Acting As",
}


def ensure_dependencies():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"], check=True)


def prompt_config():
    print("\n" + "=" * 60)
    print("  NERIS Fire Activity Report by Entity")
    print("=" * 60)
    username = input("NERIS Email: ").strip()
    print("NERIS Password (note: characters will be visible):")
    password = input("> ").strip()
    state_code = input("State Code (e.g. VA, MI, CA): ").strip().upper()
    entity_id = input("NERIS Entity ID (optional): ").strip() or None

    print("\nDate Range Options:")
    for number, option in enumerate(DATE_RANGE_OPTIONS, 1):
        print(f"  {number}. {option}")
    while True:
        choice = input("Select option [1]: ").strip() or "1"
        if choice.isdigit() and 1 <= int(choice) <= len(DATE_RANGE_OPTIONS):
            range_type = DATE_RANGE_OPTIONS[int(choice) - 1]
            break
        print("  Invalid selection, try again.")

    start_date = end_date = None
    if range_type == "Custom Date Range":
        while True:
            try:
                start_date = datetime.strptime(input("Start date (YYYY-MM-DD): ").strip(), "%Y-%m-%d").date()
                break
            except ValueError:
                print("  Invalid format, use YYYY-MM-DD.")
        while True:
            try:
                end_date = datetime.strptime(input("End date (YYYY-MM-DD): ").strip(), "%Y-%m-%d").date()
                if end_date >= start_date:
                    break
                print("  End date must be on or after the start date.")
            except ValueError:
                print("  Invalid format, use YYYY-MM-DD.")

    if not username or not password or not state_code:
        sys.exit("Email, password, and state code are required.")
    return username, password, state_code, entity_id, range_type, start_date, end_date


def calculate_date_range(range_type, start_date=None, end_date=None):
    now = datetime.now(timezone.utc)
    today = now.replace(hour=0, minute=0, second=0, microsecond=0)
    if range_type == "All Records":
        return None, None
    if range_type == "Occurred Today":
        return today, None
    if range_type == "Occurred Yesterday":
        yesterday = today - timedelta(days=1)
        return yesterday, yesterday.replace(hour=23, minute=59, second=59)
    if range_type.startswith("Occurred Last "):
        days = int(range_type.split()[2])
        return today - timedelta(days=days), None
    if range_type.startswith("Occurred in "):
        year = int(range_type.split()[-1])
        return datetime(year, 1, 1, tzinfo=timezone.utc), datetime(year, 12, 31, 23, 59, 59, tzinfo=timezone.utc)
    if range_type == "Custom Date Range":
        return (datetime.combine(start_date, datetime.min.time(), tzinfo=timezone.utc),
                datetime.combine(end_date, datetime.max.time().replace(microsecond=0), tzinfo=timezone.utc))
    return None, None


def date_range_label(range_type, start_date=None, end_date=None):
    """Return a compact, filename-safe label for the selected date range."""
    if range_type == "All Records":
        return "all-records"
    if range_type == "Custom Date Range":
        return f"{start_date:%Y-%m-%d}_to_{end_date:%Y-%m-%d}"
    if range_type.startswith("Occurred in "):
        return range_type.split()[-1]
    if range_type == "Occurred Today":
        return datetime.now().strftime("%Y-%m-%d")
    if range_type == "Occurred Yesterday":
        return (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
    if range_type.startswith("Occurred Last "):
        return f"last-{range_type.split()[2].lower()}-days"
    return "selected-range"


def authenticate(username, password):
    os.environ.update({
        "NERIS_BASE_URL": "https://api.neris.fsri.org/v1",
        "NERIS_GRANT_TYPE": "password",
        "NERIS_USERNAME": username,
        "NERIS_PASSWORD": password,
    })
    from neris_api_client import NerisApiClient
    print("\nConnecting to NERIS API...")
    client = NerisApiClient()
    print("\nCHECK YOUR EMAIL FOR THE MFA CODE")
    client.list_incidents(page_size=1)
    print("Authentication successful!")
    return client


def fetch_incidents(client, state_code, entity_id, start=None, end=None, page_size=100):
    incidents = []
    cursor = None
    cursor_only = False
    while True:
        if cursor_only:
            params = {"cursor": cursor, "page_size": page_size}
        else:
            params = {"state": state_code, "neris_id_entity": entity_id, "page_size": page_size}
            if start:
                params["call_create_start"] = start
            if end:
                params["call_create_end"] = end
            if cursor:
                params["cursor"] = cursor
        try:
            response = client.list_incidents(**params)
            status_code = getattr(response, "status_code", None)
            if status_code == 403:
                if cursor and not cursor_only:
                    print("  API rejected the continuation cursor with filters; retrying with cursor only.")
                    cursor_only = True
                    continue
                raise RuntimeError(
                    f"NERIS API returned HTTP 403 while retrieving incidents for {entity_id}. "
                    "The account may not have access to this entity or cursor."
                )
            response = response if isinstance(response, dict) else response.json()
        except Exception as error:
            status_code = getattr(getattr(error, "response", None), "status_code", None)
            if isinstance(error, RuntimeError) and str(error).startswith("NERIS API returned HTTP 403"):
                raise
            raise RuntimeError(
                f"Unable to retrieve incidents for {entity_id}"
                f" (HTTP {status_code})" if status_code else
                f"Unable to retrieve incidents for {entity_id}: {error}"
            ) from error
        batch = response.get("incidents", [])
        if not batch:
            break
        incidents.extend(batch)
        cursor = response.get("next_cursor")
        if not cursor:
            break
        cursor_only = False
    return incidents


def fetch_entities(client, state_code, entity_id):
    if entity_id:
        try:
            entity = client.get_entity(entity_id)
            return {entity_id: entity.get("name", "") if isinstance(entity, dict) else ""}
        except Exception:
            return {entity_id: ""}
    entities = {}
    page = 1
    while True:
        response = client.list_entities(state=state_code, page_size=100, page_number=page)
        response = response if isinstance(response, dict) else response.json()
        batch = response.get("entities", [])
        if not batch:
            break
        for entity in batch:
            if entity.get("neris_id"):
                entities[entity["neris_id"]] = entity.get("name", "")
        if len(batch) < 100 or len(entities) >= response.get("total_count", 0):
            break
        page += 1
    return entities


def _number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _first_number(obj, paths):
    for path in paths:
        value = obj
        for part in path.split("."):
            if not isinstance(value, dict):
                value = None
                break
            value = value.get(part)
        if value is not None and value != "":
            return _number(value)
    return 0.0


def _type_strings(incident):
    return {
        (item.get("type") or "").strip().upper()
        for item in incident.get("incident_types", []) or []
        if isinstance(item, dict) and (item.get("type") or "").strip().upper().startswith("FIRE||")
    }


def _is_fatal(injury_type):
    value = (injury_type or "").upper()
    return "FATAL" in value and "NONFATAL" not in value


def _rescue_actions(rescue):
    details = rescue.get("ffrescue_or_nonffrescue") or {}
    actions = details.get("actions") or rescue.get("actions") or []
    return {str(action).upper() for action in actions}


def _metric_values(incident):
    values = {column: 0 for column in METRIC_COLUMNS}
    exposures = incident.get("exposures") or []
    values["Number of Exposures"] = len(exposures)
    base = incident.get("base") or {}
    values["Number of Displacements"] = _number(base.get("displacement_count"))
    values["Property Value Loss"] = _first_number(incident, [
        "fire_detail.property_value_loss", "fire_detail.property_value_loss_amount",
        "fire_detail.location_detail.property_value_loss",
    ])
    values["Contents Value Loss"] = _first_number(incident, [
        "fire_detail.contents_value_loss", "fire_detail.contents_value_loss_amount",
        "fire_detail.location_detail.contents_value_loss",
    ])
    values["Acres Burned"] = _first_number(incident, [
        "fire_detail.outside_fire_acres_burned", "outside_fire_acres_burned",
        "fire_detail.location_detail.acres_burned",
    ])
    values["Total Loss"] = values["Property Value Loss"] + values["Contents Value Loss"]

    aid_columns_present = set()
    for aid in incident.get("aids") or []:
        if not isinstance(aid, dict):
            continue
        direction = str(aid.get("aid_direction") or "").upper()
        aid_type = str(aid.get("aid_type") or "").upper().replace("-", "_").replace(" ", "_")
        column = AID_COLUMNS.get((direction, aid_type))
        if column:
            aid_columns_present.add(column)
    for column in aid_columns_present:
        values[column] = 1

    for casualty_rescue in incident.get("casualty_rescues") or []:
        if not isinstance(casualty_rescue, dict):
            continue
        casualty_type = str(casualty_rescue.get("type") or "").upper()
        injury = (casualty_rescue.get("casualty") or {}).get("injury_or_noninjury") or {}
        injury_type = injury.get("type") or ""
        if casualty_type == "FF":
            prefix = "Number of FF "
            rescue_prefix = "Number of FF Rescues"
        else:
            prefix = "Number of nonFF "
            rescue_prefix = "Number of nonFF Rescues"
        if _is_fatal(injury_type):
            values[prefix + "Deaths"] += 1
        elif "NONFATAL" in str(injury_type).upper():
            values[prefix + "Injuries"] += 1
        actions = _rescue_actions(casualty_rescue.get("rescue") or {})
        if actions & {"RESCUED_BY_FIREFIGHTER", "EVAC_ASSISTED_BY_FIREFIGHTER"}:
            values[rescue_prefix] += 1
    return values


def aggregate(incidents, entity_names):
    detail = defaultdict(lambda: {column: 0 for column in METRIC_COLUMNS})
    detail_types = defaultdict(lambda: defaultdict(set))
    summary = defaultdict(lambda: {column: 0 for column in METRIC_COLUMNS})
    summary_types = defaultdict(lambda: defaultdict(set))

    for incident in incidents:
        entity_id = (incident.get("base") or {}).get("department_neris_id", "")
        if not entity_id:
            continue
        metrics = _metric_values(incident)
        incident_types = _type_strings(incident)
        for incident_type in incident_types:
            key = (entity_id, incident_type)
            detail[key]["Number of Fires"] += 1
            for column in METRIC_COLUMNS[1:]:
                if column.endswith(" Types"):
                    detail_types[key][column].update(metrics[column])
                else:
                    detail[key][column] += metrics[column]
            group = incident_type.split("||")[1] if len(incident_type.split("||")) > 1 else ""
            for label, group_code in SUMMARY_GROUPS.items():
                if group != group_code:
                    continue
                summary_key = (entity_id, label)
                summary[summary_key]["Number of Fires"] += 1
                for column in METRIC_COLUMNS[1:]:
                    if column.endswith(" Types"):
                        summary_types[summary_key][column].update(metrics[column])
                    else:
                        summary[summary_key][column] += metrics[column]

    return detail, detail_types, summary, summary_types, entity_names


def export_excel(detail, detail_types, summary, summary_types, entity_names, state_code,
                 range_label="selected-range"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(
        os.path.expanduser("~"), "Documents",
        f"neris_fire_activity_{state_code}_{range_label}_{timestamp}.xlsx",
    )
    os.makedirs(os.path.dirname(path), exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(*(Side(style="thin", color="B7C9D6") for _ in range(4)))
    heading_border = Border(*(Side(style="thin", color="000000") for _ in range(4)))
    sheet = workbook.create_sheet("Entity Fire Reports")
    headers_by_table = {
        "detail": ["Fire Incident Type"] + METRIC_COLUMNS,
        "summary": ["Fire Summary Type"] + METRIC_COLUMNS,
    }
    currency_columns = {"Property Value Loss", "Contents Value Loss", "Total Loss"}

    def append_table(data, type_sets, type_header, entity_id):
        headers = headers_by_table[type_header]
        sheet.append(headers)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for key in sorted(data):
            if key[0] != entity_id:
                continue
            _, type_name = key
            row = [type_name]
            for column in METRIC_COLUMNS:
                value = data[key][column]
                if column == "Acres Burned":
                    value = math.ceil((value - 1e-12) * 100) / 100
                row.append(value)
            sheet.append(row)
            for column_index, column in enumerate(headers, 1):
                cell = sheet.cell(sheet.max_row, column_index)
                cell.border = border
                if column in currency_columns:
                    cell.number_format = '$#,##0;($#,##0);-'
                elif column == "Acres Burned":
                    cell.number_format = '#,##0.00;(#,##0.00);-'
                elif column in METRIC_COLUMNS:
                    cell.number_format = '#,##0;(#,##0);-'

    entity_ids = sorted(set(entity_names) | {key[0] for key in detail} | {key[0] for key in summary})
    for entity_index, entity_id in enumerate(entity_ids):
        if entity_index:
            sheet.append([])
            sheet.append([])
        sheet.append([entity_id, entity_names.get(entity_id, "")])
        heading_row = sheet.max_row
        sheet.merge_cells(start_row=heading_row, start_column=2, end_row=heading_row, end_column=5)
        sheet.row_dimensions[heading_row].height = 48
        for cell in sheet[heading_row][:1]:
            cell.font = Font(bold=True, size=12)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = heading_border
        name_cell = sheet.cell(heading_row, 2)
        name_cell.font = Font(bold=True, size=12)
        name_cell.fill = PatternFill("solid", fgColor="D9EAF7")
        name_cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        name_cell.border = heading_border
        for column_index in (3, 4, 5):
            merged_cell = sheet.cell(heading_row, column_index)
            merged_cell.fill = PatternFill("solid", fgColor="D9EAF7")
            merged_cell.border = heading_border
        append_table(detail, detail_types, "detail", entity_id)
        sheet.append([])
        append_table(summary, summary_types, "summary", entity_id)

    for row_number in range(1, sheet.max_row + 1):
        if sheet.cell(row_number, 1).value in ("Fire Incident Type", "Fire Summary Type"):
            sheet.row_dimensions[row_number].height = 48

    # Header text wraps, so widths are driven by actual displayed data values.
    # This keeps numeric columns compact without truncating their contents.
    for column_index in range(1, len(METRIC_COLUMNS) + 2):
        column_letter = get_column_letter(column_index)
        longest_value = 0
        for row_number in range(1, sheet.max_row + 1):
            cell = sheet.cell(row_number, column_index)
            if cell.value is None:
                continue
            if cell.value in ("Fire Incident Type", "Fire Summary Type"):
                continue
            if column_index == 1:
                value_length = len(str(cell.value))
            elif column_index >= 2 and column_index == 2 and isinstance(cell.value, str):
                # Department names are wrapped in the light-blue heading row.
                continue
            elif column_index >= 2:
                if cell.number_format == '$#,##0;($#,##0);-':
                    value_length = len(f"${round(float(cell.value or 0)):,.0f}")
                elif cell.number_format == '#,##0.00;(#,##0.00);-':
                    value_length = len(f"{float(cell.value or 0):,.2f}")
                elif cell.number_format == '#,##0;(#,##0);-':
                    value_length = len(f"{round(float(cell.value or 0)):,}")
                else:
                    value_length = len(str(cell.value))
            else:
                value_length = 0
            longest_value = max(longest_value, value_length)
        sheet.column_dimensions[column_letter].width = max(8, min(longest_value + 2, 40))
    workbook.save(path)
    return path


def main():
    ensure_dependencies()
    username, password, state_code, entity_id, range_type, start_date, end_date = prompt_config()
    start, end = calculate_date_range(range_type, start_date, end_date)
    client = authenticate(username, password)
    entity_names = fetch_entities(client, state_code, entity_id)
    incidents = []
    for current_entity in entity_names:
        print(f"Fetching {current_entity}...")
        incidents.extend(fetch_incidents(client, state_code, current_entity, start, end))
    detail, detail_types, summary, summary_types, entity_names = aggregate(incidents, entity_names)
    if not entity_names:
        print("No departments found for the selected state/entity.")
        return
    range_label = date_range_label(range_type, start_date, end_date)
    path = export_excel(
        detail, detail_types, summary, summary_types, entity_names, state_code, range_label
    )
    print(f"\nExported: {path}")
    print(f"Incidents retrieved: {len(incidents)}")
    print(f"Fire type rows: {len(detail)}; summary rows: {len(summary)}")


if __name__ == "__main__":
    main()
