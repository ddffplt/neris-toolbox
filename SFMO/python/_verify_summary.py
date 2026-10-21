from EntityFireActivityReport import aggregate, export_excel
from openpyxl import load_workbook
import os

entity_names = {'E1': 'Alpha'}
incidents = [{
    'base': {'department_neris_id': 'E1', 'displacement_count': 2},
    'incident_types': [{'type': 'FIRE||STRUCTURE_FIRE'}],
    'exposures': [],
    'fire_detail': {'property_value_loss': 1000, 'contents_value_loss': 250},
    'aids': [],
    'casualty_rescues': []
}]

detail, detail_types, summary, summary_types, entity_names = aggregate(incidents, entity_names)
path = export_excel(detail, detail_types, summary, summary_types, entity_names, 'VA', 'test-range')
wb = load_workbook(path)
print(wb.sheetnames)
print(wb['Summary']['A1'].value)
print(wb['Summary']['C2'].value)
os.remove(path)
