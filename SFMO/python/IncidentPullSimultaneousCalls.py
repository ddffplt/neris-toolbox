"""
NERIS Incident Export Tool — Simultaneous Calls
------------------------------------
Fetches incidents from the NERIS API (statewide, county, single entity or
multiple entities) and exports only incidents that overlapped in time with at
least one other incident for the same department. An incident's window runs
from first dispatch to last unit clear (the same two times used for Total
Incident Time).

Note that if a column has no data throughout the entire dataset, it will be dropped from the Excel file output. 
"""

import sys
import subprocess
import os
import copy
import json
import re
import time
import traceback
from datetime import datetime, timedelta, timezone

def ensure_dependencies():
    def pip_install(*packages):
        subprocess.run(
            [sys.executable, "-m", "pip", "install", *packages, "--quiet"],
            check=True
        )

    try:
        import pandas   # noqa
        import openpyxl # noqa
    except ImportError:
        print("Installing pandas and openpyxl...")
        pip_install("pandas", "openpyxl")
        print("✓ pandas and openpyxl installed")

    try:
        from neris_api_client import NerisApiClient  # noqa
        print("✓ Dependencies ready")
    except ImportError:
        print("Installing neris-api-client...")
        pip_install("neris-api-client")
        print("✓ neris-api-client installed")

DATE_RANGE_OPTIONS = [
    "All Records",
    "Occurred Today",
    "Occurred Yesterday",
    "Occurred Last 7 Days",
    "Occurred Last 30 Days",
    "Occurred Last 90 Days",
    "Occurred in 2025",
    "Occurred in 2026",
    "Occurred in current calendar year to date",
    "Custom Date Range",
]

FD25_COUNTY_CODES = {f"{number:03d}" for number in range(1, 28, 2)}


def normalize_fd25_prefix(value):
    prefix = (value or "").strip().upper().rstrip("*")
    if prefix.startswith("FD25") and prefix[4:] in FD25_COUNTY_CODES:
        return prefix
    return None


def prompt_config():
    print("\n" + "=" * 60)
    print("  NERIS Incident Export Tool — Simultaneous Calls")
    print("=" * 60)

    print("\n── Credentials ──────────────────────────────────")
    username = input("NERIS Email: ").strip()
    print("NERIS Password (note: characters will be visible):")
    password = input("> ").strip()

    print("\n── Query Parameters ─────────────────────────────")
    state_code = input("State Code (e.g. MI, VA, CA): ").strip().upper()

    print("\nEntity Scope Options:")
    print("  1. Entire state")
    print("  2. County (FD25 prefix)")
    print("  3. Single entity")
    print("  4. Multiple entities")
    while True:
        scope_choice = input("Select entity scope [1]: ").strip() or "1"
        if scope_choice in {"1", "2", "3", "4"}:
            break
        print("  Invalid selection, choose 1, 2, 3, or 4.")

    entity_id = None
    entity_prefix = None
    entity_names = []
    if scope_choice == "2":
        while True:
            entity_prefix = normalize_fd25_prefix(
                input("FD25 county prefix (FD25001*, FD25003*, ... FD25027*): ")
            )
            if entity_prefix:
                break
            print("  Invalid prefix. Use FD25 followed by an odd county code from 001 through 027.")
    elif scope_choice == "3":
        entity_id = input("NERIS Entity ID or exact entity name: ").strip()
        if not entity_id:
            print("  An entity ID or name is required for this scope.")
            return prompt_config()
    elif scope_choice == "4":
        entity_names = [
            value.strip()
            for value in input("Entity IDs or exact names, separated by commas: ").split(",")
            if value.strip()
        ]
        if not entity_names:
            print("  At least one entity is required.")
            return prompt_config()

    entity_scope = {
        "1": "statewide",
        "2": "fd25_prefix",
        "3": "entity",
        "4": "multiple_entities",
    }[scope_choice]

    print("\nDate Range Options:")
    for i, opt in enumerate(DATE_RANGE_OPTIONS, 1):
        print(f"  {i}. {opt}")
    while True:
        choice = input("Select option [1]: ").strip() or "1"
        if choice.isdigit() and 1 <= int(choice) <= len(DATE_RANGE_OPTIONS):
            range_type = DATE_RANGE_OPTIONS[int(choice) - 1]
            break
        print("  Invalid selection, try again.")

    start_date = end_date = None
    if range_type == "Custom Date Range":
        while True:
            s = input("Start date (YYYY-MM-DD): ").strip()
            try:
                start_date = datetime.strptime(s, "%Y-%m-%d").date()
                break
            except ValueError:
                print("  Invalid format, use YYYY-MM-DD.")
        while True:
            e = input("End date (YYYY-MM-DD): ").strip()
            try:
                end_date = datetime.strptime(e, "%Y-%m-%d").date()
                if end_date < start_date:
                    print("  End date must be on or after the start date.")
                    continue
                break
            except ValueError:
                print("  Invalid format, use YYYY-MM-DD.")

    return (
        username, password, state_code, entity_id, entity_prefix, entity_names,
        entity_scope, range_type, start_date, end_date,
    )

def calculate_date_range(range_type, start_date=None, end_date=None):
    """Return (call_create_start, call_create_end) as UTC-aware datetimes."""
    now   = datetime.now(tz=timezone.utc)
    today = now.replace(hour=0, minute=0, second=0, microsecond=0)

    if range_type == "All Records":
        return None, None
    elif range_type == "Occurred Today":
        return today, None
    elif range_type == "Occurred Yesterday":
        yesterday = today - timedelta(days=1)
        return yesterday, yesterday.replace(hour=23, minute=59, second=59)
    elif range_type == "Occurred Last 7 Days":
        return today - timedelta(days=7), None
    elif range_type == "Occurred Last 30 Days":
        return today - timedelta(days=30), None
    elif range_type == "Occurred Last 90 Days":
        return today - timedelta(days=90), None
    elif range_type == "Occurred in current calendar year to date":
        return datetime(now.year, 1, 1, tzinfo=timezone.utc), now
    elif range_type.startswith("Occurred in "):
        year = int(range_type.split()[-1])
        return (
            datetime(year, 1,  1,  0,  0,  0, tzinfo=timezone.utc),
            datetime(year, 12, 31, 23, 59, 59, tzinfo=timezone.utc),
        )
    elif range_type == "Custom Date Range" and start_date and end_date:
        start_dt = datetime.combine(start_date, datetime.min.time(), tzinfo=timezone.utc)
        end_dt   = datetime.combine(
            end_date, datetime.max.time().replace(microsecond=0), tzinfo=timezone.utc
        )
        return start_dt, end_dt

    return None, None


def authenticate(username, password):
    from neris_api_client import NerisApiClient, Config

    print("\nConnecting to NERIS API...")
    client = NerisApiClient(Config(
        base_url="https://api.neris.fsri.org/v1",
        grant_type="password",
        username=username,
        password=password,
    ))

    print("\n✓ Authentication successful!")
    return client


def response_as_dict(response):
    if isinstance(response, dict):
        return response

    status_code = getattr(response, "status_code", None)
    if status_code is not None and status_code >= 400:
        url = getattr(response, "url", "")
        endpoint = url.split("?", 1)[0] if url else "the requested endpoint"
        detail = str(getattr(response, "text", "") or "").strip()
        if len(detail) > 300:
            detail = detail[:300] + "..."
        detail_suffix = f" Response: {detail}" if detail else ""
        raise RuntimeError(
            f"NERIS API returned HTTP {status_code} for {endpoint}."
            f"{detail_suffix} A 403 may reflect entity permissions or an API gateway restriction."
        )

    try:
        return response.json()
    except (AttributeError, ValueError) as exc:
        raise RuntimeError("NERIS API returned a response that was not valid JSON.") from exc


def list_incidents_page(client, params, cursor, page_size):
    for attempt in range(3):
        try:
            response = client.list_incidents(**params)
            status_code = getattr(response, "status_code", None)
            if status_code in (401, 403) and cursor:
                cursor_only_params = {"cursor": cursor, "page_size": page_size}
                if params != cursor_only_params:
                    print("  API rejected the continuation cursor with filters; retrying with cursor only.")
                    response = client.list_incidents(**cursor_only_params)
            return response_as_dict(response)
        except Exception:
            if attempt == 2:
                raise
            delay = 2 ** attempt
            print(f"  Request failed; retrying in {delay} second(s) ({attempt + 1}/2)...")
            time.sleep(delay)


def incident_call_create(incident):
    value = (incident.get("dispatch") or {}).get("call_create")
    if isinstance(value, datetime):
        parsed = value
    elif value:
        try:
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except ValueError:
            return None
    else:
        return None
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed


def incident_dedupe_key(incident):
    incident_id = incident.get("neris_id")
    if incident_id:
        return str(incident_id)
    return json.dumps(incident, sort_keys=True, default=str)


def get_state_incidents(client, state_code,
                        call_create_start=None, call_create_end=None,
                        page_size=100):
    all_incidents = []
    next_cursor   = None
    page_count    = 0

    print(f"\nFetching incidents for state: {state_code}")
    if call_create_start:
        print(f"  call_create >= {call_create_start:%Y-%m-%d %H:%M:%S}")
    if call_create_end:
        print(f"  call_create <= {call_create_end:%Y-%m-%d %H:%M:%S}")
    if not call_create_start and not call_create_end:
        print("  No date filter — retrieving all records")

    while True:
        page_count += 1
        print(f"  Page {page_count}... ", end="", flush=True)

        kwargs = {"state": state_code, "page_size": page_size}
        if call_create_start: kwargs["call_create_start"] = call_create_start
        if call_create_end:   kwargs["call_create_end"]   = call_create_end
        if next_cursor:       kwargs["cursor"]             = next_cursor

        try:
            res = list_incidents_page(client, kwargs, next_cursor, page_size)
        except Exception as e:
            print(f"\n✗ Error on page {page_count}: {e}")
            traceback.print_exc()
            break

        incidents = res.get("incidents", [])
        if not incidents:
            print("empty page — done.")
            break

        all_incidents.extend(incidents)
        print(f"{len(incidents)} fetched  (running total: {len(all_incidents)})")

        next_cursor = res.get("next_cursor")
        if not next_cursor:
            print("  No more pages.")
            break

    print(f"\n{'='*50}")
    print(f"Total incidents retrieved: {len(all_incidents)}")
    print(f"{'='*50}")
    return all_incidents


def get_entity_incidents(client, entity_id,
                         call_create_start=None, call_create_end=None,
                         page_size=100):
    all_incidents = []
    seen_incidents = set()
    next_cursor = None
    page_count = 0
    active_end = call_create_end
    last_recovery_end = None

    print(f"\nFetching incidents for entity: {entity_id}")
    while True:
        page_count += 1
        print(f"  Page {page_count}... ", end="", flush=True)
        kwargs = {"neris_id_entity": entity_id, "page_size": page_size}
        if call_create_start:
            kwargs["call_create_start"] = call_create_start
        if active_end:
            kwargs["call_create_end"] = active_end
        if next_cursor:
            kwargs["cursor"] = next_cursor

        try:
            res = list_incidents_page(client, kwargs, next_cursor, page_size)
        except Exception as e:
            print(f"\n✗ Error on page {page_count}: {e}")
            traceback.print_exc()
            known_times = [
                parsed for parsed in (incident_call_create(item) for item in all_incidents)
                if parsed is not None
            ]
            if not known_times:
                print("  Cannot resume this entity because no fetched incident has a usable call_create time.")
                break

            recovery_end = min(known_times)
            if active_end and recovery_end > active_end:
                recovery_end = active_end
            if last_recovery_end and recovery_end >= last_recovery_end:
                recovery_end = last_recovery_end - timedelta(microseconds=1)
                print(
                    "  No earlier incident was reached; moving the recovery cutoff just before "
                    f"{last_recovery_end.isoformat()}. Incidents at that exact timestamp may be missed."
                )
            if call_create_start and recovery_end <= call_create_start:
                print("  The recovery cutoff reached the requested start time; moving to the next entity.")
                break
            if active_end and recovery_end >= active_end:
                print("  Could not move the recovery cutoff earlier; moving to the next entity.")
                break

            print(
                "  Resuming this entity without the failed cursor, up to "
                f"{recovery_end.isoformat()} (previously fetched incidents will be deduplicated)."
            )
            active_end = recovery_end
            last_recovery_end = recovery_end
            next_cursor = None
            continue

        incidents = res.get("incidents", [])
        if not incidents:
            print("empty page — done.")
            break

        new_incidents = []
        for incident in incidents:
            key = incident_dedupe_key(incident)
            if key not in seen_incidents:
                seen_incidents.add(key)
                all_incidents.append(incident)
                new_incidents.append(incident)
        print(f"{len(new_incidents)} fetched  (running total: {len(all_incidents)})")
        next_cursor = res.get("next_cursor")
        if not next_cursor:
            print("  No more pages.")
            break

    return all_incidents


def normalize_entity_name(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def get_scoped_incidents(client, state_code, entity_id=None, entity_prefix=None,
                         entity_names=None, entity_scope="statewide",
                         call_create_start=None, call_create_end=None,
                         page_size=100):
    if entity_scope == "entity" and entity_id and entity_id.upper().startswith("FD"):
        return get_entity_incidents(
            client, entity_id,
            call_create_start=call_create_start,
            call_create_end=call_create_end,
            page_size=page_size,
        )

    scope_label = f"entities beginning with {entity_prefix}" if entity_prefix else f"state {state_code}"
    print(f"\nFetching incidents for {scope_label} by enumerating entities...")
    entities = []
    page_number = 1
    while True:
        try:
            res = client.list_entities(
                state=state_code,
                page_size=page_size,
                page_number=page_number,
            )
            res = response_as_dict(res)
        except Exception as e:
            print(f"\n✗ Could not enumerate entities: {e}")
            traceback.print_exc()
            return []

        batch = res.get("entities", [])
        if not batch:
            break
        entities.extend(batch)
        print(f"  Retrieved {len(batch)} entities (total so far: {len(entities)})")
        total_count = res.get("total_count", 0)
        if (total_count and len(entities) >= total_count) or len(batch) < page_size:
            break
        page_number += 1

    if entity_prefix:
        entities = [
            entity for entity in entities
            if str(entity.get("neris_id", "")).startswith(entity_prefix)
        ]
    elif entity_scope in {"entity", "multiple_entities"}:
        requested = [entity_id] if entity_scope == "entity" else (entity_names or [])
        selected_entities = []
        missing = []
        ambiguous = []
        for value in requested:
            normalized_value = value.strip()
            wanted = normalize_entity_name(normalized_value)

            def base_name(entity):
                name = normalize_entity_name(entity.get("name", ""))
                for suffix in ("firedepartment", "firedept", "fire"):
                    if name.endswith(suffix):
                        return name[:-len(suffix)]
                return name

            tiers = [
                lambda e: str(e.get("neris_id", "")).upper() == normalized_value.upper(),
                lambda e: bool(wanted) and normalize_entity_name(e.get("name", "")) == wanted,
                lambda e: bool(wanted) and base_name(e) == wanted,
                lambda e: bool(wanted) and wanted in normalize_entity_name(e.get("name", "")),
            ]
            matching = []
            for tier in tiers:
                matching = [entity for entity in entities if tier(entity)]
                if matching:
                    break
            if not matching:
                missing.append(normalized_value)
            elif len(matching) > 1:
                ambiguous.append(
                    f"{normalized_value} ({', '.join(str(item.get('name', '')) for item in matching)})"
                )
            else:
                selected_entities.append(matching[0])
        if missing or ambiguous:
            details = []
            if missing:
                details.append("could not resolve: " + ", ".join(missing))
            if ambiguous:
                details.append("ambiguous: " + "; ".join(ambiguous))
            raise ValueError("Entity selection failed: " + " | ".join(details))
        entities = list(dict.fromkeys(
            (entity.get("neris_id"), entity.get("name"))
            for entity in selected_entities
        ))
        entities = [
            {"neris_id": entity_id_value, "name": entity_name}
            for entity_id_value, entity_name in entities
        ]
    print(f"\nFound {len(entities)} entities to query")

    all_incidents = []
    for entity in entities:
        current_entity_id = entity.get("neris_id") if isinstance(entity, dict) else None
        if not current_entity_id:
            continue
        print(f"\nQuerying entity {current_entity_id} ({entity.get('name', '')}) ...")
        all_incidents.extend(
            get_entity_incidents(
                client,
                current_entity_id,
                call_create_start=call_create_start,
                call_create_end=call_create_end,
                page_size=page_size,
            )
        )

    print(f"\nTotal aggregated incidents: {len(all_incidents)}")
    return all_incidents


def parse_api_datetime(value):
    if value in (None, ""):
        return None
    try:
        if isinstance(value, (int, float)):
            return datetime.fromtimestamp(value, tz=timezone.utc)
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (TypeError, ValueError, OverflowError):
        return None
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed


def to_eastern_text(value):
    """Convert a UTC-aware datetime to US Eastern (EST/EDT) as 'YYYY-MM-DD HH:MM:SS'.

    DST runs from the second Sunday in March 2:00 local (07:00 UTC) to the first
    Sunday in November 2:00 local (06:00 UTC); implemented directly to avoid a
    tzdata dependency on Windows.
    """
    utc = value.astimezone(timezone.utc).replace(tzinfo=None)

    def nth_sunday(month, n):
        first = datetime(utc.year, month, 1)
        return first + timedelta(days=(6 - first.weekday()) % 7 + 7 * (n - 1))

    dst_start = nth_sunday(3, 2).replace(hour=7)
    dst_end = nth_sunday(11, 1).replace(hour=6)
    offset = 4 if dst_start <= utc < dst_end else 5
    return (utc - timedelta(hours=offset)).strftime("%Y-%m-%d %H:%M:%S")


def incident_window(incident):
    """Start = earliest unit dispatch (fallback: call dispatch); end = last unit clear.

    These are the same two times used for Total Incident Time.
    """
    units = [u for u in (incident.get("unit_responses") or []) if isinstance(u, dict)]
    starts = [t for t in (parse_api_datetime(u.get("dispatch")) for u in units) if t]
    if not starts:
        fallback = parse_api_datetime((incident.get("dispatch") or {}).get("call_dispatch"))
        starts = [fallback] if fallback else []
    ends = [t for t in (parse_api_datetime(u.get("unit_clear")) for u in units) if t]
    if not starts or not ends or max(ends) < min(starts):
        return None
    return min(starts), max(ends)


def find_simultaneous(incidents):
    """Return (incidents with an overlap, {neris_id: info}) per department.

    Two incidents are simultaneous when they start at the same time or their
    windows overlap by more than zero time (one starting exactly as another
    ends is not simultaneous).
    """
    by_dept = {}
    skipped = 0
    for inc in incidents:
        window = incident_window(inc)
        if not window:
            skipped += 1
            continue
        dept_id = (inc.get("base") or {}).get("department_neris_id", "")
        by_dept.setdefault(dept_id, []).append((window[0], window[1], inc))

    info = {}
    for items in by_dept.values():
        for start, end, inc in items:
            others = [
                other for o_start, o_end, other in items
                if other is not inc
                and (o_start == start or (o_start < end and start < o_end))
            ]
            if others:
                info[inc.get("neris_id")] = {
                    "start": start,
                    "end": end,
                    "count": len(others) + 1,
                    "others": [
                        str((o.get("base") or {}).get("incident_number") or o.get("neris_id", ""))
                        for o in others
                    ],
                }
    if skipped:
        print(f"  {skipped} incident(s) skipped: no usable dispatch/clear times.")
    simultaneous = [inc for inc in incidents if inc.get("neris_id") in info]
    return simultaneous, info


def export_to_excel(incidents, client, state_code, entity_scope="statewide",
                    entity_id=None, entity_prefix=None, entity_names=None,
                    simultaneous_info=None):
    from report_metadata import add_report_sources_sheet, incident_date_range
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    if not incidents:
        print("⚠ No incidents to export.")
        return None

    date_str = datetime.now().strftime("%Y-%m-%d")
    time_str = datetime.now().strftime("%H-%M-%S")
    filename = f"NERIS Simultaneous Calls {state_code} {date_str} {time_str}.xlsx"

    # Build department name cache for all unique dept IDs in the results
    dept_ids = {
        (inc.get("base") or {}).get("department_neris_id")
        for inc in incidents
    } - {None, ""}

    print(f"\nLooking up {len(dept_ids)} department names...")
    dept_names = {}
    for i, did in enumerate(dept_ids, 1):
        try:
            entity = client.get_entity(did)
            dept_names[did] = entity.get("name", "")
        except Exception:
            dept_names[did] = ""
        if i % 25 == 0:
            print(f"  {i}/{len(dept_ids)} departments looked up...")
    print("✓ Department lookup complete")

    # ── Helpers ───────────────────────────────────────────────────────────────

    def extract_list_field(lst, key):
        if not lst or not isinstance(lst, list):
            return ""
        return "||".join(
            str(item.get(key, ""))
            for item in lst
            if isinstance(item, dict) and item.get(key)
        )

    def join_list(lst):
        if not lst or not isinstance(lst, list):
            return ""
        return "||".join(str(x) for x in lst if x)

    def parse_api_datetime(value):
        if value in (None, ""):
            return None
        try:
            if isinstance(value, (int, float)):
                return datetime.fromtimestamp(value, tz=timezone.utc)
            return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except (TypeError, ValueError, OverflowError):
            return None

    simultaneous_info = simultaneous_info or {}

    def simultaneous_columns(inc):
        sim = simultaneous_info.get(inc.get("neris_id"), {})
        return {
            "Incident Start (Eastern Time Zone)": to_eastern_text(sim["start"]) if sim else "",
            "Incident End (Eastern Time Zone)": to_eastern_text(sim["end"]) if sim else "",
            "Simultaneous Call Count": sim.get("count", ""),
            "Simultaneous Incident Numbers": " || ".join(sim.get("others", [])),
        }

    def summarize_units(units, call_dispatch=None):
        # Time on scene: first unit on scene to last unit clear.
        # Total incident time: first dispatch (earliest unit dispatch, falling
        # back to the call dispatch time) to last unit clear.
        valid = [u for u in units if isinstance(u, dict)]
        arrivals = [t for t in (parse_api_datetime(u.get("on_scene")) for u in valid) if t]
        clears = [t for t in (parse_api_datetime(u.get("unit_clear")) for u in valid) if t]
        dispatches = [t for t in (parse_api_datetime(u.get("dispatch")) for u in valid) if t]
        if not dispatches:
            fallback = parse_api_datetime(call_dispatch)
            dispatches = [fallback] if fallback else []

        def span_minutes(starts, ends):
            if not starts or not ends:
                return ""
            duration = (max(ends) - min(starts)).total_seconds() / 60
            return round(duration, 2) if duration >= 0 else ""

        return {
            "total_units": len(valid),
            "time_on_scene": span_minutes(arrivals, clears),
            "total_incident_time": span_minutes(dispatches, clears),
        }

    def get_coordinates_from_point(point_obj):
        if not point_obj or not isinstance(point_obj, dict):
            return None, None
        geometry = point_obj.get("geometry", {})
        if geometry and isinstance(geometry, dict):
            coords = geometry.get("coordinates", [])
            if coords and len(coords) >= 2:
                return coords[1], coords[0]   # GeoJSON is [lon, lat]
        return None, None

    def get_location_data(loc_dict, point_obj=None):
        if not loc_dict or not isinstance(loc_dict, dict):
            return {}
        lat, lon = get_coordinates_from_point(point_obj)
        return {
            "street_number": loc_dict.get("number", ""),
            "street_prefix": loc_dict.get("street_prefix_direction", ""),
            "street_name":   loc_dict.get("street", ""),
            "street_type":   loc_dict.get("street_postfix", ""),
            "street_suffix": loc_dict.get("street_postfix_direction", ""),
            "apartment":     loc_dict.get("unit_value", ""),
            "city":          loc_dict.get("incorporated_municipality", ""),
            "state":         loc_dict.get("state", ""),
            "zip_code":      loc_dict.get("postal_code", ""),
            "county":        loc_dict.get("county", ""),
            "latitude":      lat if lat is not None else "",
            "longitude":     lon if lon is not None else "",
            "census_tract":  "",
            "census_block":  "",
            "parcel_id":     "",
        }

    ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

    def sanitize(val):
        return ILLEGAL_CHARS_RE.sub("", val) if isinstance(val, str) else val

    # ── Overview sheet ────────────────────────────────────────────────────────

    overview_rows = []
    for inc in incidents:
        base       = inc.get("base", {}) or {}
        dispatch   = inc.get("dispatch", {}) or {}
        status_obj = inc.get("incident_status", {}) or {}
        dept_id    = base.get("department_neris_id", "")

        inc_types = inc.get("incident_types", []) or []
        type_list = [
            it.get("type", "")
            for it in inc_types
            if isinstance(it, dict) and it.get("type")
        ]

        unit_summary = summarize_units(
            inc.get("unit_responses", []) or [], dispatch.get("call_dispatch")
        )

        overview_rows.append({
            "Incident NERIS ID":    inc.get("neris_id", ""),
            "Department/Entity ID": dept_id,
            "Department Name":      dept_names.get(dept_id, ""),
            "Incident Number":      base.get("incident_number", ""),
            "Call Date":            dispatch.get("call_create", ""),
            "Last Modified":        inc.get("last_modified", ""),
            "Incident Status":      status_obj.get("status", ""),
            "First Incident Type":  type_list[0] if len(type_list) > 0 else "",
            "Second Incident Type": type_list[1] if len(type_list) > 1 else "",
            "Third Incident Type":  type_list[2] if len(type_list) > 2 else "",
            "People Present":       base.get("people_present", ""),
            "Outcome Narrative":    base.get("outcome_narrative", ""),
            "Displacement Count":   base.get("displacement_count", ""),
            "Total Units":          unit_summary["total_units"],
            "Incident Time on Scene (Minutes)": unit_summary["time_on_scene"],
            "Total Incident Time (Minutes)": unit_summary["total_incident_time"],
            **simultaneous_columns(inc),
            "Submitter Type":       inc.get("submitter_account_type", ""),
        })

    # ── Raw Data ──────────────────────────────────────────────────────────────

    processed = []

    for inc in copy.deepcopy(incidents):
        base     = inc.get("base", {}) or {}
        dispatch = inc.get("dispatch", {}) or {}
        dept_id  = base.get("department_neris_id", "")

        inc["department_name"]      = dept_names.get(dept_id, "")
        inc["department_neris_id"]  = dept_id
        inc["incident_number"]      = base.get("incident_number", "")
        inc["department_time_zone"] = (inc.get("department", {}) or {}).get("time_zone", "")

        inc_types = inc.get("incident_types", []) or []
        inc["primary_incident_types"] = "||".join(
            it.get("type", "") for it in inc_types
            if isinstance(it, dict) and it.get("primary")
        )
        inc["other_incident_types"] = "||".join(
            it.get("type", "") for it in inc_types
            if isinstance(it, dict) and not it.get("primary")
        )
        inc.pop("incident_types", None)

        inc["special_modifiers"] = extract_list_field(inc.get("special_modifiers", []), "type")

        aids = inc.get("aids", []) or []
        inc["aid_types"]          = extract_list_field(aids, "aid_type")
        inc["aid_directions"]     = extract_list_field(aids, "aid_direction")
        inc["aid_department_ids"] = extract_list_field(aids, "department_neris_id")
        inc.pop("aids", None)

        inc["nonfd_aid_types"] = extract_list_field(inc.get("nonfd_aids", []), "type")
        inc.pop("nonfd_aids", None)

        units = inc.get("unit_responses", []) or []
        summary = summarize_units(units, dispatch.get("call_dispatch"))
        inc["unit_ids"]                      = extract_list_field(units, "reported_unit_id")
        inc["unit_neris_ids"]                = extract_list_field(units, "unit_neris_id")
        inc["total_units"]                   = summary["total_units"]
        inc["total_staffing"]                = sum(
            u.get("staffing", 0) or 0 for u in units if isinstance(u, dict)
        )
        inc["unit_response_modes"]           = extract_list_field(units, "response_mode")
        inc["incident_time_on_scene_minutes"] = summary["time_on_scene"]
        inc["total_incident_time_minutes"]    = summary["total_incident_time"]
        for col, val in simultaneous_columns(inc).items():
            inc[re.sub(r"[^a-z0-9]+", "_", col.lower()).strip("_")] = val
        inc.pop("unit_responses", None)

        disp_units = dispatch.get("unit_responses", []) or []
        dispatch["disp_unit_ids"]       = extract_list_field(disp_units, "reported_unit_id")
        dispatch["disp_total_units"]    = len(disp_units)
        dispatch["disp_total_staffing"] = sum(
            u.get("staffing", 0) or 0 for u in disp_units if isinstance(u, dict)
        )
        dispatch.pop("unit_responses", None)

        comments = dispatch.get("comments", []) or []
        dispatch["comments_text"] = extract_list_field(comments, "comment")
        dispatch.pop("comments", None)

        base["displacement_causes"] = join_list(base.get("displacement_causes", []))

        # Location consolidation
        base_loc       = base.get("location", {}) or {}
        base_geo       = base_loc.get("geocoded_location", {}) or {}
        base_geo_point = base_geo.get("point", {}) if base_geo else {}
        disp_loc       = dispatch.get("location", {}) or {}
        disp_geo       = disp_loc.get("geocoded_location", {}) or {}
        disp_geo_point = disp_geo.get("point", {}) if disp_geo else {}

        location_source = ""
        final_location  = {}

        if base_geo and any(base_geo.values()):
            final_location  = get_location_data(base_geo, base_geo_point)
            location_source = "base_location_geocoded"
        elif base_loc and any(base_loc.values()):
            final_location  = get_location_data(base_loc)
            location_source = "base_location"
        elif disp_geo and any(disp_geo.values()):
            final_location  = get_location_data(disp_geo, disp_geo_point)
            location_source = "dispatch_location_geocoded"
        elif disp_loc and any(disp_loc.values()):
            final_location  = get_location_data(disp_loc)
            location_source = "dispatch_location"

        inc["location_source"] = location_source
        for field, value in final_location.items():
            inc[f"location_{field}"] = value
        base.pop("location", None)
        dispatch.pop("location", None)

        parcel = inc.get('parcel', {}) or {}
        parcel_assessment = (parcel.get('summary', {}) or {}).get('assessment', {}) or {}
        inc['parcel_primary_year_built'] = parcel_assessment.get('primary_year_built', '')
        inc['parcel_primary_zoning'] = parcel_assessment.get('primary_zoning', '')
        inc['parcel_primary_use'] = parcel_assessment.get('primary_use', '')
        inc['parcel_primary_owner_occupied'] = parcel_assessment.get('primary_owner_occupied', '')
        inc['parcel_primary_assessment_land'] = parcel_assessment.get('primary_assessment_land', '')
        inc['parcel_computed_loss'] = parcel_assessment.get('computed_loss', '')
        inc.pop('parcel', None)
        
        exposures = inc.get("exposures", []) or []
        inc["exposure_count"]        = len(exposures)
        inc["exposure_damage_types"] = extract_list_field(exposures, "damage_type")
        inc.pop("exposures", None)

        inc.pop("casualty_rescues", None)

        med = inc.get("medical_details", []) or []
        inc["med_report_ids"]             = extract_list_field(med, "patient_care_report_id")
        inc["med_evaluations"]            = extract_list_field(med, "patient_care_evaluation")
        inc["med_patient_statuses"]       = extract_list_field(med, "patient_status")
        inc["med_transport_dispositions"] = extract_list_field(med, "transport_disposition")
        inc.pop("medical_details", None)

        elec = inc.get("electric_hazards", []) or []
        inc["electric_hazard_types"]         = extract_list_field(elec, "type")
        inc["electric_hazard_source_target"] = extract_list_field(elec, "source_or_target")
        inc.pop("electric_hazards", None)

        inc["powergen_hazard_count"] = len(inc.get("powergen_hazards") or [])
        inc.pop("powergen_hazards", None)

        fire = inc.get("fire_detail", {}) or {}
        fire["investigation_types"]    = join_list(fire.get("investigation_types", []))
        fire["suppression_appliances"] = join_list(fire.get("suppression_appliances", []))

        hazsit = inc.get("hazsit_detail", {}) or {}
        chems  = hazsit.get("chemicals", []) or []
        hazsit["chemical_names"]       = extract_list_field(chems, "name")
        hazsit["chemical_dot_classes"] = extract_list_field(chems, "dot_class")
        hazsit.pop("chemicals", None)

        for alarm_key in ("smoke_alarm", "fire_alarm", "other_alarm"):
            presence = (inc.get(alarm_key, {}) or {}).get("presence", {}) or {}
            presence["alarm_types"] = join_list(presence.get("alarm_types", []))

        fs_presence = (inc.get("fire_suppression", {}) or {}).get("presence", {}) or {}
        supp_types  = fs_presence.get("suppression_types", []) or []
        fs_presence["suppression_types"] = (
            extract_list_field(supp_types, "type")
            if supp_types and isinstance(supp_types[0], dict)
            else join_list(supp_types)
        )

        cfs_presence = (inc.get("cooking_fire_suppression", {}) or {}).get("presence", {}) or {}
        cfs_presence["suppression_types"] = join_list(cfs_presence.get("suppression_types", []))

        at = inc.get("actions_tactics", {}) or {}
        an = at.get("action_noaction", {}) or {}
        an["actions"] = join_list(an.get("actions", []))

        processed.append(inc)

    df_raw = pd.json_normalize(processed, sep="_")

  # If you want to add any of these fields back, you can grab them from this list and add them to the lists above. 
    exclude_cols = {
        "actions_tactics_last_modified", "actions_tactics_neris_uid",
        "base_location_geocoded_location_neris_uid", "base_location_neris_uid",
        "base_location_use_in_use_last_modified", "base_location_use_in_use_neris_uid",
        "base_location_use_neris_uid",
        "census_tract_incident_neris_id", "census_tract_last_modified", "census_tract_neris_uid",
        "cooking_fire_suppression_last_modified", "cooking_fire_suppression_neris_uid",
        "cooking_fire_suppression_presence_last_modified", "cooking_fire_suppression_presence_neris_uid",
        "csst_hazard_last_modified", "csst_hazard_neris_uid",
        "dispatch_last_modified", "dispatch_location_geocoded_location_neris_uid", "dispatch_neris_uid",
        "dispatch_tactic_timestamps_last_modified", "dispatch_tactic_timestamps_neris_uid",
        "fire_alarm_last_modified", "fire_alarm_neris_uid",
        "fire_alarm_presence_last_modified", "fire_alarm_presence_neris_uid",
        "fire_detail_last_modified", "fire_detail_neris_uid",
        "fire_suppression_last_modified", "fire_suppression_neris_uid",
        "fire_suppression_presence_last_modified", "fire_suppression_presence_neris_uid",
        "fire_suppression_presence_operation_type_effectiveness_last_modified",
        "fire_suppression_presence_operation_type_effectiveness_neris_uid",
        "fire_suppression_presence_operation_type_last_modified",
        "fire_suppression_presence_operation_type_neris_uid",
        "hazsit_detail_last_modified", "hazsit_detail_neris_uid",
        "incident_status_last_modified",
        "other_alarm_last_modified", "other_alarm_neris_uid",
        "other_alarm_presence_last_modified", "other_alarm_presence_neris_uid",
        "smoke_alarm_last_modified", "smoke_alarm_neris_uid",
        "smoke_alarm_presence_operation_alerted_failed_other_last_modified",
        "smoke_alarm_presence_operation_alerted_failed_other_neris_uid",
        "smoke_alarm_presence_operation_last_modified", "smoke_alarm_presence_operation_neris_uid",
        "tactic_timestamps_last_modified", "tactic_timestamps_neris_uid",
        "weather_last_modified", "weather_neris_uid",
        "department_time_zone", "base_neris_uid", "base_last_modified",
        "base_department_neris_id", "incident_status_created_by",
        "dispatch_location_neris_uid",
        "tactic_timestamps_command_established", "tactic_timestamps_completed_sizeup",
        "tactic_timestamps_suppression_complete", "tactic_timestamps_primary_search_begin",
        "tactic_timestamps_primary_search_complete", "tactic_timestamps_water_on_fire",
        "tactic_timestamps_fire_under_control", "tactic_timestamps_fire_knocked_down",
        "weather_incident_neris_id",
        "dispatch_tactic_timestamps_command_established",
        "dispatch_tactic_timestamps_completed_sizeup",
        "dispatch_tactic_timestamps_suppression_complete",
        "dispatch_tactic_timestamps_fire_under_control",
    }
    df_raw = df_raw.drop(columns=[c for c in exclude_cols if c in df_raw.columns])

    def reorder_columns(df):
        first_cols = ["department_neris_id", "department_name", "neris_id",
                      "neris_uid", "incident_status_status"]
        location_cols = [
            "location_source", "location_street_number", "location_street_prefix",
            "location_street_name", "location_street_type", "location_street_suffix",
            "location_apartment", "location_city", "location_state",
            "location_zip_code", "location_county", "location_latitude",
            "location_longitude", "location_census_tract", "location_census_block",
            "location_parcel_id",
        ]
        hazard_cols = ["electric_hazard_types", "electric_hazard_source_target",
                       "powergen_hazard_count"]

        existing_first    = [c for c in first_cols    if c in df.columns]
        existing_location = [c for c in location_cols if c in df.columns]
        existing_hazard   = [c for c in hazard_cols   if c in df.columns]
        pinned            = set(existing_first + existing_location + existing_hazard + ["last_modified"])
        middle_cols       = [c for c in df.columns if c not in pinned]

        final_order = existing_first + existing_location + middle_cols + existing_hazard
        if "last_modified" in df.columns:
            final_order.append("last_modified")
        return df[final_order]

    df_raw = reorder_columns(df_raw)

    def clean_df(df):
        drop = [
            col for col in df.columns
            if df[col].dropna().pipe(
                lambda s: s.empty or s.apply(
                    lambda x: x == "" or x == [] or str(x) == "[]"
                ).all()
            )
        ]
        return df.drop(columns=drop)

    df_overview   = clean_df(pd.DataFrame(overview_rows).fillna(""))
    df_raw        = clean_df(df_raw)

    print(f"  Overview:          {len(df_overview.columns)} columns")
    print(f"  Raw Data:          {len(df_raw.columns)} columns")

    def write_sheet(wb, df, name):
        ws      = wb.create_sheet(name)
        headers = list(df.columns)
        eastern_col_idx = None
        utc_col_idx = None
        if "Call Date" in headers:
            utc_col_idx = headers.index("Call Date") + 1
            headers.insert(utc_col_idx, "Call Date (Eastern Time Zone)")
            eastern_col_idx = utc_col_idx + 1
            headers[utc_col_idx + 1:utc_col_idx + 1] = [
                "Call Date (Eastern Time Zone) Date",
                "Call Date (Eastern Time Zone) Time",
                "Call Date (Eastern Time Zone) Month",
                "Call Date (Eastern Time Zone) Year",
                "Call Date (Eastern Time Zone) Hour",
            ]
        ws.append(headers)
        for _, row in df.iterrows():
            values = [sanitize(row[h]) if h in row else "" for h in headers]
            ws.append(values)
            if eastern_col_idx is not None:
                row_num = ws.max_row
                utc_cell = f"{get_column_letter(utc_col_idx)}{row_num}"
                eastern_cell = f"{get_column_letter(eastern_col_idx)}{row_num}"
                ws.cell(row_num, eastern_col_idx).value = (
                    f'=IF({utc_cell}="","",'
                    f'DATE(VALUE(LEFT({utc_cell},4)),VALUE(MID({utc_cell},6,2)),'
                    f'VALUE(MID({utc_cell},9,2)))+'
                    f'TIME(VALUE(MID({utc_cell},12,2)),VALUE(MID({utc_cell},15,2)),'
                    f'VALUE(MID({utc_cell},18,2)))-IF('
                    f'AND(DATE(VALUE(LEFT({utc_cell},4)),VALUE(MID({utc_cell},6,2)),'
                    f'VALUE(MID({utc_cell},9,2)))+TIME(VALUE(MID({utc_cell},12,2)),'
                    f'VALUE(MID({utc_cell},15,2)),VALUE(MID({utc_cell},18,2)))>='
                    f'DATE(VALUE(LEFT({utc_cell},4)),3,1)+'
                    f'MOD(8-WEEKDAY(DATE(VALUE(LEFT({utc_cell},4)),3,1)),7)+7+TIME(7,0,0),'
                    f'DATE(VALUE(LEFT({utc_cell},4)),3,1)+'
                    f'MOD(8-WEEKDAY(DATE(VALUE(LEFT({utc_cell},4)),3,1)),7)+7+TIME(7,0,0)<'
                    f'DATE(VALUE(LEFT({utc_cell},4)),11,1)+'
                    f'MOD(8-WEEKDAY(DATE(VALUE(LEFT({utc_cell},4)),11,1)),7)+TIME(6,0,0)), '
                    f'TIME(4,0,0),TIME(5,0,0)))'
                )
                ws.cell(row_num, eastern_col_idx + 1).value = (
                    f'=IF({eastern_cell}="","",INT({eastern_cell}))'
                )
                ws.cell(row_num, eastern_col_idx + 2).value = (
                    f'=IF({eastern_cell}="","",MOD({eastern_cell},1))'
                )
                ws.cell(row_num, eastern_col_idx + 3).value = (
                    f'=IF({eastern_cell}="","",CHOOSE(MONTH({eastern_cell}),'
                    f'"Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"))'
                )
                ws.cell(row_num, eastern_col_idx + 4).value = (
                    f'=IF({eastern_cell}="","",YEAR({eastern_cell}))'
                )
                ws.cell(row_num, eastern_col_idx + 5).value = (
                    f'=IF({eastern_cell}="","",HOUR({eastern_cell}))'
                )
                ws.cell(row_num, eastern_col_idx).number_format = "mm/dd/yy hh:mm"
                ws.cell(row_num, eastern_col_idx + 1).number_format = "mm/dd/yy"
                ws.cell(row_num, eastern_col_idx + 2).number_format = "hh:mm"


        fill  = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        font  = Font(color="FFFFFF", bold=True, size=11)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cell in ws[1]:
            cell.fill, cell.font, cell.alignment = fill, font, align

        for col_idx, header in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            max_len    = len(str(header))
            for row_num in range(2, min(ws.max_row + 1, 52)):
                val = ws.cell(row_num, col_idx).value
                if val:
                    max_len = max(max_len, len(str(val)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)
        ws.freeze_panes = "A2"

    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.remove(wb.active)
    write_sheet(wb, df_overview, "Overview")
    write_sheet(wb, df_raw,      "Raw Data")

    import pathlib
    docs_dir = pathlib.Path.home() / "Documents"
    docs_dir.mkdir(exist_ok=True)
    filename = str(docs_dir / filename)
    sheets = "Overview, Raw Data"
    source_details = [
        ("State", state_code),
        ("Entity scope", {
            "entity": "Individual entity",
            "fd25_prefix": "FD25 county prefix",
            "statewide": "Entire state",
            "multiple_entities": "Multiple entities",
        }.get(entity_scope, entity_scope)),
        ("Entity ID", entity_id or ""),
    ]
    if entity_scope == "multiple_entities":
        name_to_id = {
            re.sub(r"[^a-z0-9]", "", str(n).casefold()): d
            for d, n in dept_names.items() if n
        }
        resolved_ids = []
        for value in entity_names or []:
            value = str(value).strip()
            key = re.sub(r"[^a-z0-9]", "", value.casefold())
            resolved_ids.append(
                value if value.upper() in {str(d).upper() for d in dept_names}
                else name_to_id.get(key, value)
            )
        resolved_ids = list(dict.fromkeys(resolved_ids))
        source_details.append(("Entity ID(s)", " || ".join(resolved_ids)))
        source_details.append((
            "Department Name",
            " || ".join(dept_names.get(d, "") for d in resolved_ids if dept_names.get(d)),
        ))
    elif entity_scope == "entity":
        source_details.append(("Department Name", dept_names.get(entity_id, "")))
    elif entity_scope == "fd25_prefix":
        county_department_names = sorted({
            name for name in dept_names.values() if name
        })
        source_details.append((
            "Department Name",
            " || ".join(county_department_names),
        ))
    source_details.extend([
        ("County prefix", entity_prefix or ""),
        ("Incident types", "All (simultaneous calls only)"),
        ("Included sheets", sheets),
    ])

    add_report_sources_sheet(
        wb,
        "IncidentPullSimultaneousCalls.py",
        incident_date_range(incidents),
        source_details,
    )
    wb.save(filename)

    print(f"\n✓ Exported: {filename}")
    print(f"  Incidents : {len(incidents)}")
    print(f"  Sheets    : {sheets}")
    return filename

def main():
    ensure_dependencies()

    (
        username, password, state_code, entity_id, entity_prefix, entity_names,
        entity_scope, range_type, start_date, end_date,
    ) = prompt_config()

    if not username or not password:
        sys.exit("✗ Email and password are required.")
    if not state_code:
        sys.exit("✗ State code is required.")

    client = authenticate(username, password)

    start_filter, end_filter = calculate_date_range(range_type, start_date, end_date)

    incidents = get_scoped_incidents(
        client,
        state_code,
        entity_id=entity_id,
        entity_prefix=entity_prefix,
        entity_names=entity_names,
        entity_scope=entity_scope,
        call_create_start=start_filter,
        call_create_end=end_filter,
    )

    total_before_filter = len(incidents)
    incidents, simultaneous_info = find_simultaneous(incidents)
    print(
        f"\nSimultaneous incidents: {len(incidents)} of {total_before_filter} "
        "overlap with another incident from the same department."
    )
    if not incidents:
        print("⚠ No simultaneous calls found for the given parameters.")
        return

    print("\nExporting to Excel...")
    export_to_excel(
        incidents,
        client,
        state_code=state_code,
        entity_scope=entity_scope,
        entity_id=entity_id,
        entity_prefix=entity_prefix,
        entity_names=entity_names,
        simultaneous_info=simultaneous_info,
    )

    print("\n" + "=" * 60)
    print("  PROCESS COMPLETE")
    print("=" * 60)


if __name__ == "__main__":
    main()
